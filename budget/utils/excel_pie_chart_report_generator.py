import os

from openpyxl import Workbook
from openpyxl.chart import PieChart3D, Reference

from .CONF import generate_report_filename


def generate_excel_with_pie_chart(income_report_data, expense_report_data):
    """
    Формат report_data должен быть следующим
        data = [
        ("Type of Expense", "Amount Spent"),
        ("Grocery", 300),
        ("Electricity", 150),
    ]
    """
    wb_pie_chart = Workbook()
    ws_expense = wb_pie_chart.active
    ws_expense.title = "Expenses"

    for row in expense_report_data:
        ws_expense.append(row)
    expense_pie = PieChart3D()
    labels = Reference(ws_expense, min_col=1, min_row=2, max_row=7)
    data = Reference(ws_expense, min_col=2, min_row=1, max_row=7)
    expense_pie.add_data(data, titles_from_data=True)
    expense_pie.set_categories(labels)
    expense_pie.title = "Expenditures Pie Chart"
    ws_expense.add_chart(expense_pie, "C10")

    # Создаем круговую диаграмму для доходов
    ws_income = wb_pie_chart.create_sheet(title="Incomes")
    for row in income_report_data:
        ws_income.append(row)
    income_pie = PieChart3D()
    labels = Reference(ws_income, min_col=1, min_row=2, max_row=len(income_report_data))
    data = Reference(ws_income, min_col=2, min_row=1, max_row=len(income_report_data))
    income_pie.add_data(data, titles_from_data=True)
    income_pie.set_categories(labels)
    income_pie.title = "Income Pie Chart"
    ws_income.add_chart(income_pie, "C10")

    output_dir = "reports"
    os.makedirs(output_dir, exist_ok=True)
    filepath = os.path.join(output_dir, generate_report_filename("pie_chart_report"))
    wb_pie_chart.save(filepath)
    print(f"Файл успешно сохранен: {os.path.abspath(filepath)}")
